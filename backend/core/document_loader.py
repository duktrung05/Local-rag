import re
import os
import csv
from pathlib import Path
from typing import List, Dict, Any
import logging

logger = logging.getLogger(__name__)


class DocumentLoader:
    """Load và làm sạch nội dung tài liệu từ nhiều định dạng"""
    
    SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".docx", ".md", ".csv", ".xlsx", ".xls"}
    
    def load_file(self, file_path: str) -> List[Dict[str, Any]]:
        """
        Load một file và trả về list các document chunks với metadata
        
        Returns:
            List[Dict]: [{"content": str, "metadata": dict}, ...]
        """
        path = Path(file_path)
        
        if not path.exists():
            raise FileNotFoundError(f"File không tồn tại: {file_path}")
        
        ext = path.suffix.lower()
        
        if ext not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(f"Định dạng không hỗ trợ: {ext}")
        
        logger.info(f"Loading file: {path.name} ({ext})")
        
        loaders = {
            ".pdf": self._load_pdf,
            ".txt": self._load_text,
            ".md": self._load_text,
            ".docx": self._load_docx,
            ".csv": self._load_csv,
            ".xlsx": self._load_excel,
            ".xls": self._load_excel,
        }
        
        raw_docs = loaders[ext](str(path))
        
        # Gán metadata cho từng document thô
        loaded_docs = []
        for doc in raw_docs:
            if doc["content"].strip():  # Bỏ qua doc rỗng
                doc["metadata"]["file_name"] = path.name
                doc["metadata"]["file_path"] = str(path)
                doc["metadata"]["file_type"] = ext
                loaded_docs.append(doc)
        
        logger.info(f"Loaded {len(loaded_docs)} documents từ {path.name}")
        return loaded_docs
    
    def load_directory(self, dir_path: str) -> List[Dict[str, Any]]:
        """Load tất cả file được hỗ trợ trong thư mục"""
        dir_path = Path(dir_path)
        all_docs = []
        
        for file_path in dir_path.rglob("*"):
            if file_path.suffix.lower() in self.SUPPORTED_EXTENSIONS:
                try:
                    docs = self.load_file(str(file_path))
                    all_docs.extend(docs)
                except Exception as e:
                    logger.error(f"Lỗi khi load {file_path}: {e}")
        
        return all_docs
    
    def _load_pdf(self, file_path: str) -> List[Dict[str, Any]]:
        """Load file PDF"""
        try:
            from pypdf import PdfReader
            
            reader = PdfReader(file_path)
            docs = []
            
            for i, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                if text.strip():
                    docs.append({
                        "content": text,
                        "metadata": {
                            "page": i + 1,
                            "total_pages": len(reader.pages),
                            "source": file_path
                        }
                    })
            
            return docs
        except Exception as e:
            logger.error(f"Lỗi load PDF: {e}")
            raise
    
    def _load_text(self, file_path: str) -> List[Dict[str, Any]]:
        """Load file TXT hoặc MD"""
        try:
            with open(file_path, "r", encoding="utf-8-sig", errors="ignore") as f:
                content = f.read()
            
            return [{
                "content": content,
                "metadata": {"source": file_path}
            }]
        except Exception as e:
            logger.error(f"Lỗi load text file: {e}")
            raise
    
    def _load_docx(self, file_path: str) -> List[Dict[str, Any]]:
        """Load file DOCX"""
        try:
            from docx import Document
            
            doc = Document(file_path)
            paragraphs = []
            
            for para in doc.paragraphs:
                if para.text.strip():
                    paragraphs.append(para.text)
            
            # Gộp thành các section
            content = "\n\n".join(paragraphs)
            
            return [{
                "content": content,
                "metadata": {"source": file_path}
            }]
        except Exception as e:
            logger.error(f"Lỗi load DOCX: {e}")
            raise
    
    def _load_csv(self, file_path: str) -> List[Dict[str, Any]]:
        """Load file CSV - mỗi row là một document"""
        try:
            docs = []
            
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                reader = csv.DictReader(f)
                headers = reader.fieldnames or []
                
                for i, row in enumerate(reader):
                    # Convert row to readable text
                    content = " | ".join([f"{k}: {v}" for k, v in row.items() if v])
                    
                    if content.strip():
                        docs.append({
                            "content": content,
                            "metadata": {
                                "row": i + 1,
                                "source": file_path,
                                "columns": ", ".join(headers)
                            }
                        })
            
            return docs
        except Exception as e:
            logger.error(f"Lỗi load CSV: {e}")
            raise
    
    def _load_excel(self, file_path: str) -> List[Dict[str, Any]]:
        """Load file Excel (.xlsx, .xls)"""
        try:
            import openpyxl
            from pathlib import Path
            
            suffix = Path(file_path).suffix.lower()
            if suffix == ".xls":
                # openpyxl does not support legacy .xls files
                raise ValueError(
                    "Định dạng Excel cũ (.xls) không được openpyxl hỗ trợ trực tiếp. "
                    "Vui lòng chuyển đổi tệp thành định dạng Excel mới (.xlsx) và thử lại."
                )
                
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            docs = []
            
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                rows = list(sheet.iter_rows(values_only=True))
                if not rows:
                    continue
                
                # First non-empty row as header
                header_row_idx = 0
                while header_row_idx < len(rows) and not any(rows[header_row_idx]):
                    header_row_idx += 1
                    
                if header_row_idx >= len(rows):
                    continue
                    
                headers = [str(h).strip() if h is not None else f"Column_{j+1}" 
                           for j, h in enumerate(rows[header_row_idx])]
                
                for i, row_vals in enumerate(rows[header_row_idx + 1:]):
                    row_dict = {}
                    for col_idx, val in enumerate(row_vals):
                        if val is not None and str(val).strip():
                            header = headers[col_idx] if col_idx < len(headers) else f"Column_{col_idx+1}"
                            row_dict[header] = str(val).strip()
                    
                    if row_dict:
                        content = " | ".join([f"{k}: {v}" for k, v in row_dict.items()])
                        if content.strip():
                            docs.append({
                                "content": content,
                                "metadata": {
                                    "sheet": sheet_name,
                                    "row": i + header_row_idx + 2,
                                    "source": file_path,
                                    "columns": ", ".join(headers[:20]) # Limit columns metadata size
                                }
                            })
            return docs
        except Exception as e:
            logger.error(f"Lỗi load Excel: {e}")
            raise

    

